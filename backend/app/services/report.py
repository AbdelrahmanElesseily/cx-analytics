from __future__ import annotations
import io
from collections import defaultdict
from datetime import date as _date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

_GOLD   = '#B8860B'
_BLUE   = '#2563EB'
_GREEN  = '#059669'
_PURPLE = '#7C3AED'
_AMBER  = '#D97706'
_SLATE  = '#94A3B8'

_STAGE_COLORS = {
    'Service Application and Submission': _BLUE,
    'Communication During Procedures':    _GOLD,
    'Receiving Service Information':       _GREEN,
    'Service Completion':                  _PURPLE,
}

plt.rcParams.update({
    'font.family': 'sans-serif',
    'axes.spines.top': False,
    'axes.spines.right': False,
    'figure.facecolor': 'white',
    'axes.facecolor': 'white',
})


def _save(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=130, bbox_inches='tight')
    plt.close(fig)
    return buf.getvalue()


def _generate_charts(feedback: list, improvements: list) -> list[tuple[str, bytes]]:
    charts: list[tuple[str, bytes]] = []

    ch_map: dict = defaultdict(lambda: {'s': 0, 'n': 0})
    for r in feedback:
        ch = r.get('channel') or 'Unknown'
        ch_map[ch]['s'] += r.get('rating', 0)
        ch_map[ch]['n'] += 1
    if ch_map:
        data = sorted([(c, v['s']/v['n']) for c, v in ch_map.items()], key=lambda x: x[1])
        names, vals = zip(*data)
        fig, ax = plt.subplots(figsize=(7, max(2.5, len(names) * 0.45)))
        ax.barh(names, vals, color=_GOLD, alpha=0.85)
        ax.set_xlim(0, 5); ax.set_title('Avg Satisfaction by Channel', fontweight='bold', fontsize=11)
        ax.tick_params(labelsize=8)
        charts.append(('Avg Satisfaction by Channel', _save(fig)))

    q_map: dict = defaultdict(int)
    for r in improvements:
        q_map[r.get('quarter') or '?'] += 1
    if q_map:
        qs, counts = zip(*sorted(q_map.items()))
        fig, ax = plt.subplots(figsize=(7, 3.2))
        ax.bar(qs, counts, color=_BLUE, alpha=0.85)
        ax.set_title('Improvement Actions by Quarter', fontweight='bold', fontsize=11)
        ax.tick_params(labelsize=9)
        charts.append(('Improvement Actions by Quarter', _save(fig)))

    stage_map: dict = defaultdict(int)
    for r in improvements:
        stage_map[r.get('cx_stage') or 'Unknown'] += 1
    if stage_map:
        labels = list(stage_map.keys())
        sizes  = list(stage_map.values())
        colors = [_STAGE_COLORS.get(l, _SLATE) for l in labels]
        fig, ax = plt.subplots(figsize=(7, 4))
        wedges, _, autotexts = ax.pie(sizes, colors=colors, autopct='%1.0f%%', startangle=90,
                                       pctdistance=0.82, wedgeprops={'width': 0.55})
        for t in autotexts: t.set_fontsize(8)
        ax.legend(wedges, labels, loc='center left', bbox_to_anchor=(1, 0.5), fontsize=8)
        ax.set_title('Issues by CX Stage', fontweight='bold', fontsize=11)
        charts.append(('Issues by CX Stage', _save(fig)))

    rating_map: dict = defaultdict(int)
    for r in feedback:
        rating_map[r.get('rating', 0)] += 1
    if rating_map:
        ratings = sorted(rating_map.keys())
        fig, ax = plt.subplots(figsize=(4.5, 3.2))
        ax.bar([f'★{r}' for r in ratings], [rating_map[r] for r in ratings], color=_GREEN, alpha=0.85)
        ax.set_title('Feedback Rating Distribution', fontweight='bold', fontsize=11)
        ax.tick_params(labelsize=9)
        charts.append(('Feedback Rating Distribution', _save(fig)))

    chan_map: dict = defaultdict(int)
    for r in improvements:
        chan_map[r.get('channel') or 'Unknown'] += 1
    if chan_map:
        data_c = sorted(chan_map.items(), key=lambda x: x[1])[-8:]
        names, counts = zip(*data_c)
        fig, ax = plt.subplots(figsize=(7, max(2.5, len(names) * 0.45)))
        ax.barh(names, counts, color=_AMBER, alpha=0.85)
        ax.set_title('Improvement Actions by Channel', fontweight='bold', fontsize=11)
        ax.tick_params(labelsize=8)
        charts.append(('Improvement Actions by Channel', _save(fig)))

    src_map: dict = defaultdict(int)
    for r in improvements:
        src = (r.get('source') or 'Unknown').replace('Employee S&F & Contact Center Suggestion', 'Emp+CC')
        src_map[src] += 1
    if src_map:
        data_s = sorted(src_map.items(), key=lambda x: x[1])
        names, counts = zip(*data_s)
        fig, ax = plt.subplots(figsize=(7, max(2.5, len(names) * 0.45)))
        ax.barh(names, counts, color=_GREEN, alpha=0.85)
        ax.set_title('Improvement Actions by Source', fontweight='bold', fontsize=11)
        ax.tick_params(labelsize=8)
        charts.append(('Improvement Actions by Source', _save(fig)))

    return charts


def build_report(
    sections: list[dict],
    cover: dict | None = None,
    include_charts: bool = True,
    feedback: list | None = None,
    improvements: list | None = None,
) -> bytes:
    doc  = Document()
    cov  = cover or {}
    today = _date.today().strftime('%B %d, %Y')

    # ── Cover page ─────────────────────────────────────────────────────────
    t = doc.add_heading(cov.get('title') or 'CX Analytics Report', level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _muted(text):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        p.runs[0].font.size = Pt(10)

    if cov.get('author'):      _muted(cov['author'])
    if cov.get('department'):  _muted(cov['department'])
    _muted(cov.get('date') or today)
    doc.add_paragraph('')

    # ── Dashboard charts ────────────────────────────────────────────────────
    if include_charts and (feedback or improvements):
        doc.add_heading('Dashboard Overview', level=1)
        for title_txt, png in _generate_charts(feedback or [], improvements or []):
            doc.add_heading(title_txt, level=3)
            doc.add_picture(io.BytesIO(png), width=Inches(5.8))
            doc.add_paragraph('')
        doc.add_page_break()

    # ── Sections ────────────────────────────────────────────────────────────
    findings = [s for s in sections if s.get('type') == 'finding']
    if findings:
        doc.add_heading('AI Assistant Findings', level=1)
        doc.add_paragraph('')

    finding_idx = 0
    for sec in sections:
        stype = sec.get('type', 'finding')

        if stype == 'custom':
            sec_title = sec.get('sectionTitle') or 'Notes'
            doc.add_heading(sec_title, level=2)
            if sec.get('content'):
                doc.add_paragraph(sec['content'])
            doc.add_paragraph('')
            continue

        if stype == 'chart':
            chart_title = sec.get('chartTitle', 'Dashboard Chart')
            doc.add_heading(chart_title, level=2)
            rows = sec.get('data', [])
            data_key = sec.get('dataKey', '')
            label_key = sec.get('labelKey', '')
            chart_type = sec.get('chartType', 'bar')
            if rows and data_key and label_key:
                names  = [str(r.get(label_key, ''))[:22] for r in rows]
                values = [float(r.get(data_key, 0)) for r in rows]
                horizontal = chart_type in ('hbar',) or len(names) > 5
                fig, ax = plt.subplots(figsize=(7, max(2.8, len(names) * 0.42)))
                colors = [_BLUE, _GOLD, _GREEN, _PURPLE, _AMBER, _GREEN, _SLATE, _BLUE]
                bar_colors = [colors[i % len(colors)] for i in range(len(names))]
                if horizontal:
                    ax.barh(names, values, color=bar_colors, alpha=0.85)
                else:
                    ax.bar(names, values, color=bar_colors, alpha=0.85)
                    ax.tick_params(axis='x', labelrotation=30, labelsize=8)
                ax.set_title(chart_title, fontweight='bold', fontsize=11)
                ax.tick_params(labelsize=8)
                doc.add_picture(io.BytesIO(_save(fig)), width=Inches(5.8))
            doc.add_paragraph('')
            continue

        # finding
        finding_idx += 1
        question = sec.get('question', '')
        mode     = sec.get('mode', '')
        data     = sec.get('data', {})
        note     = sec.get('note', '')

        doc.add_heading(f'Q{finding_idx}: {question}', level=2)

        if note:
            p = doc.add_paragraph(note)
            p.runs[0].font.italic = True
            p.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        if mode == 'rag':
            answer = data.get('answer') or ''
            if answer:
                doc.add_paragraph(answer)
            sources = data.get('sources', [])
            if sources:
                doc.add_heading('Sources', level=3)
                for s in sources[:3]:
                    parts = [s.get('type',''), s.get('quarter',''), s.get('channel','')]
                    label = '  ·  '.join(p for p in parts if p)
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(label).bold = True
                    snippet = s.get('snippet', '')
                    if snippet:
                        p.add_run(f':  {snippet}')

        elif mode == 'nlq':
            insights = data.get('insights') or ''
            if insights:
                doc.add_paragraph(insights)
            rows = data.get('rows', [])
            cols = data.get('columns', [])
            if rows and cols:
                # inline chart if data is chartable
                num_cols = [c for c in cols if all(
                    r.get(c) is not None and str(r.get(c,'')).replace('.','',1).lstrip('-').isdigit()
                    for r in rows[:20]
                )]
                cat_cols = [c for c in cols if c not in num_cols]
                if num_cols and cat_cols:
                    label_key, value_key = cat_cols[0], num_cols[0]
                    chart_rows = rows[:8]
                    names  = [str(r.get(label_key, ''))[:18] for r in chart_rows]
                    values = [float(r.get(value_key, 0)) for r in chart_rows]
                    fig, ax = plt.subplots(figsize=(6, max(2.5, len(names) * 0.38)))
                    colors = [_BLUE, _GOLD, _GREEN, _PURPLE, _AMBER, _GREEN, _SLATE, _BLUE]
                    ax.barh(names, values, color=[colors[i % len(colors)] for i in range(len(names))], alpha=0.85)
                    ax.set_title(f'{value_key} by {label_key}', fontweight='bold', fontsize=10)
                    ax.tick_params(labelsize=8)
                    doc.add_picture(io.BytesIO(_save(fig)), width=Inches(5.5))
                    doc.add_paragraph('')

                display_rows = rows[:30]
                table = doc.add_table(rows=1 + len(display_rows), cols=len(cols))
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for j, col in enumerate(cols):
                    hdr[j].text = col
                    hdr[j].paragraphs[0].runs[0].bold = True
                for ri, row in enumerate(display_rows, 1):
                    cells = table.rows[ri].cells
                    for j, col in enumerate(cols):
                        cells[j].text = str(row.get(col, ''))
                if data.get('row_count', 0) > 30:
                    ex = doc.add_paragraph(f'  …and {data["row_count"] - 30} more rows')
                    ex.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                    ex.runs[0].font.size = Pt(9)

        doc.add_paragraph('')

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
